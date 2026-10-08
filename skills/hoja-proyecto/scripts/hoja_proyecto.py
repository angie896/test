#!/usr/bin/env python3
"""
hoja_proyecto.py — la hoja maestra de un proyecto (Excel / Google Sheets).

Modo sin conector (archivo). Con el conector de Google Sheets, Claude hace estas mismas
operaciones directo sobre la hoja; ver SKILL.md.

    python hoja_proyecto.py crear  PROYECTO.xlsx --titulo "NOMBRE" [--tipo serie]
    python hoja_proyecto.py agregar PROYECTO.xlsx desglose.json [--propuesta propuesta.json]
    python hoja_proyecto.py prompt PROYECTO.xlsx ID "nuevo prompt" --cambio "qué cambió"

- crear:    hoja vacía con LÉEME, DESGLOSE GENERAL e HISTORIAL PROMPTS.
- agregar:  suma un desglose (un episodio, una secuencia, un corto…). Si el ítem ya existe
            (misma categoría + mismo nombre), NO lo duplica: agrega la unidad y las notas nuevas.
            Con --propuesta, llena PROMPT y lo registra en el historial.
- prompt:   reemplaza el prompt vigente de un ID. El anterior queda en HISTORIAL PROMPTS.
Nunca toca REFERENCE ni FINAL (las imágenes las pone Angie).
"""
import argparse
import datetime as dt
import json

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

MONO = "Courier New"
HEAD = PatternFill("solid", fgColor="3F3F3F")
BANDS = {"SETS": "E4E9F0", "CHARACTERS": "F2F2F2", "PROPS": "E9EEE6",
         "GRAPHIC DESIGN": "F3ECE4", "VEHICLES": "EFEFEF", "MAKEUP / SFX": "F4E9EE", "VFX": "ECE9F3"}
PROMPT_FILL = PatternFill("solid", fgColor="FFF8E1")
THIN = Side(style="thin", color="A6A6A6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="top", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
FLAG = "C55A11"

COLS = [  # (encabezado, ancho)
    ("ID", 10), ("CATEGORY", 15), ("ITEM", 28), ("UNIDADES", 14), ("ESC.", 10),
    ("DESCRIPTION", 60), ("REFERENCE", 24), ("FINAL", 24), ("STATUS", 13),
    ("PROMPT", 70), ("VERSIÓN", 9), ("HERRAMIENTA", 16), ("NOTAS / AJUSTES", 32),
]
C = {name: i + 1 for i, (name, _) in enumerate(COLS)}
HIST = [("FECHA", 12), ("ID", 10), ("ITEM", 26), ("VERSIÓN", 9), ("PROMPT", 80),
        ("QUÉ CAMBIÓ", 34), ("¿FUNCIONÓ?", 12)]
STATUSES = ["TO DO", "EN PROCESO", "REVISIÓN", "APROBADO", "DESCARTADO"]
PREFIX = {"SETS": "SET", "CHARACTERS": "CHR", "PROPS": "PRP", "GRAPHIC DESIGN": "GD",
          "VEHICLES": "VEH", "MAKEUP / SFX": "MUP", "VFX": "VFX"}
G, H = "DESGLOSE GENERAL", "HISTORIAL PROMPTS"


def header(ws, cols, row=1):
    for i, (name, width) in enumerate(cols, start=1):
        c = ws.cell(row, i, name)
        c.fill, c.border, c.alignment = HEAD, BORDER, CENTER
        c.font = Font(name=MONO, bold=True, color="FFFFFF", size=10)
        ws.column_dimensions[c.column_letter].width = width
    ws.row_dimensions[row].height = 24
    ws.freeze_panes = ws.cell(row + 1, 4 if cols is COLS else 1)
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def crear(path, titulo, tipo):
    wb = Workbook()
    readme = wb.active
    readme.title = "LÉEME"
    readme.sheet_view.showGridLines = False
    readme.column_dimensions["A"].width = 110
    lines = [
        (f"{titulo.upper()} — DESGLOSE GENERAL DE ARTE", Font(name=MONO, bold=True, size=16)),
        (f"Tipo de proyecto: {tipo}", Font(name=MONO, size=10, color="595959")),
        ("", None),
        ("CÓMO FUNCIONA", Font(name=MONO, bold=True, size=11)),
        ("• DESGLOSE GENERAL: una fila por cada cosa del proyecto (set, personaje, prop, gráfica). "
         "Si aparece en varios episodios o secuencias, es la MISMA fila: la columna UNIDADES dice dónde aparece.", None),
        ("• Cada fila tiene un ID fijo (SET-001, CHR-004…). Para pedirle algo a Claude basta con decir el ID: "
         "\"ajusta el prompt de SET-003, más cálido\".", None),
        ("• PROMPT es el prompt vigente, listo para copiar. Cuando se ajusta, el anterior NO se pierde: "
         "queda en HISTORIAL PROMPTS con la fecha y qué cambió.", None),
        ("• REFERENCE y FINAL son tuyas: pega ahí las imágenes. Claude nunca las toca.", None),
        ("• En naranja: lo que el guion no dice y tienes que definir tú.", Font(name="Arial", size=10, color=FLAG, bold=True)),
        ("• STATUS: TO DO → EN PROCESO → REVISIÓN → APROBADO (o DESCARTADO).", None),
        ("• En HISTORIAL PROMPTS puedes marcar ¿FUNCIONÓ? (sí / no) y así sabes qué prompts sirven.", None),
    ]
    for r, (txt, font) in enumerate(lines, start=1):
        c = readme.cell(r, 1, txt)
        c.font = font or Font(name="Arial", size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet(G)
    header(ws, COLS)
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    L = ws.cell(1, C["STATUS"]).column_letter
    dv.add(f"{L}2:{L}2000")

    hs = wb.create_sheet(H)
    header(hs, HIST)
    dv2 = DataValidation(type="list", formula1='"sí,no,a medias"', allow_blank=True)
    hs.add_data_validation(dv2)
    L2 = hs.cell(1, 7).column_letter
    dv2.add(f"{L2}2:{L2}5000")
    wb.active = 1
    wb.save(path)
    print(f"OK → {path} creado")


def style_row(ws, r, category, description):
    fill = PatternFill("solid", fgColor=BANDS.get(category, "FFFFFF"))
    for name, col in C.items():
        c = ws.cell(r, col)
        c.border = BORDER
        c.font = Font(name=MONO, size=9, bold=name in ("ID", "ITEM"))
        c.alignment = CENTER if name in ("ID", "CATEGORY", "ITEM", "UNIDADES", "ESC.", "STATUS", "VERSIÓN") else WRAP
        if name == "PROMPT":
            c.fill = PROMPT_FILL
        elif name not in ("REFERENCE", "FINAL"):
            c.fill = fill
    if description and "⚠" in description:
        ws.cell(r, C["DESCRIPTION"]).font = Font(name=MONO, size=9, color=FLAG)
    lines = sum(1 + len(l) // 70 for l in (description or "").split("\n"))
    ws.row_dimensions[r].height = max(80, min(400, lines * 12 + 10))


def index(ws):
    idx, counters = {}, {}
    for r in range(2, ws.max_row + 1):
        rid, cat, item = (ws.cell(r, C[k]).value for k in ("ID", "CATEGORY", "ITEM"))
        if not rid:
            continue
        idx[(str(cat).upper(), str(item).strip().lower())] = r
        p, n = str(rid).split(".")[0].rsplit("-", 1)
        counters[p] = max(counters.get(p, 0), int(n))
    return idx, counters


def log(hs, rid, item, version, prompt, cambio):
    r = hs.max_row + 1
    vals = [dt.date.today().isoformat(), rid, item, version, prompt, cambio, None]
    for i, v in enumerate(vals, start=1):
        c = hs.cell(r, i, v)
        c.border, c.alignment = BORDER, WRAP if i in (5, 6) else CENTER
        c.font = Font(name=MONO, size=9)
    hs.row_dimensions[r].height = 90


def merge_text(old, new, unit):
    if not new or (old and new in old):
        return old
    return f"{old}\n[{unit}] {new}" if old else new


def agregar(path, desglose, propuesta):
    wb = load_workbook(path)
    ws, hs = wb[G], wb[H]
    with open(desglose, encoding="utf-8") as f:
        d = json.load(f)
    prompts = {}
    if propuesta:
        with open(propuesta, encoding="utf-8") as f:
            for it in json.load(f)["items"]:
                if it.get("prompts"):
                    prompts[(it["category"].upper(), it["item"].strip().lower())] = it
    idx, counters = index(ws)
    nuevos = actualizados = con_prompt = 0
    for sh in d["sheets"]:
        unit = sh["name"]
        for it in sh["items"]:
            cat = it["category"].upper()
            key = (cat, it["item"].strip().lower())
            desc = it.get("description", "")
            if key in idx:
                r = idx[key]
                u = ws.cell(r, C["UNIDADES"])
                if unit not in str(u.value or "").split(", "):
                    u.value = f"{u.value}, {unit}" if u.value else unit
                e = ws.cell(r, C["ESC."])
                entry = f'{unit}: {it["scenes"]}' if it.get("scenes") else ""
                if entry and entry not in str(e.value or "").split("\n"):
                    e.value = f"{e.value}\n{entry}" if e.value else entry
                dc = ws.cell(r, C["DESCRIPTION"])
                dc.value = merge_text(dc.value, desc, unit)
                actualizados += 1
            else:
                p = PREFIX.get(cat, "OTR")
                counters[p] = counters.get(p, 0) + 1
                r = ws.max_row + 1 if ws.cell(ws.max_row, 1).value else ws.max_row
                r = max(r, 2)
                vals = {"ID": f"{p}-{counters[p]:03d}", "CATEGORY": cat, "ITEM": it["item"],
                        "UNIDADES": unit,
                        "ESC.": f'{unit}: {it["scenes"]}' if it.get("scenes") else "", "DESCRIPTION": desc,
                        "STATUS": "TO DO"}
                for k, v in vals.items():
                    ws.cell(r, C[k], v)
                idx[key] = r
                nuevos += 1
            style_row(ws, r, cat, ws.cell(r, C["DESCRIPTION"]).value)
            if key in prompts and not ws.cell(r, C["PROMPT"]).value:
                base_id = ws.cell(r, C["ID"]).value
                for n, pr in enumerate(prompts[key]["prompts"]):
                    if n == 0:
                        row = r
                    else:  # variante: fila propia justo debajo, ID con sufijo .2, .3…
                        row = r + n
                        ws.insert_rows(row)
                        idx, counters = index(ws)
                        ws.cell(row, C["ID"], f"{base_id}.{n + 1}")
                        ws.cell(row, C["CATEGORY"], cat)
                        ws.cell(row, C["ITEM"], f'{it["item"]} — {pr.get("label", "")}')
                        ws.cell(row, C["UNIDADES"], unit)
                        ws.cell(row, C["STATUS"], "TO DO")
                    ws.cell(row, C["PROMPT"], pr["text"])
                    ws.cell(row, C["VERSIÓN"], "v1")
                    ws.cell(row, C["HERRAMIENTA"], pr.get("tool", ""))
                    if n == 0 and pr.get("label"):
                        ws.cell(row, C["NOTAS / AJUSTES"], pr["label"])
                    log(hs, ws.cell(row, C["ID"]).value, ws.cell(row, C["ITEM"]).value, "v1",
                        pr["text"], "versión inicial")
                    style_row(ws, row, cat, ws.cell(row, C["DESCRIPTION"]).value)
                    con_prompt += 1
    for r in range(2, ws.max_row + 1):  # insert_rows no mueve alturas ni estilos: re-aplicar
        if ws.cell(r, C["ID"]).value:
            style_row(ws, r, str(ws.cell(r, C["CATEGORY"]).value), ws.cell(r, C["DESCRIPTION"]).value)
    wb.save(path)
    print(f"OK → {path}: {nuevos} nuevos, {actualizados} actualizados, {con_prompt} prompts")


def nuevo_prompt(path, rid, text, cambio):
    wb = load_workbook(path)
    ws, hs = wb[G], wb[H]
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, C["ID"]).value == rid:
            v = str(ws.cell(r, C["VERSIÓN"]).value or "v0").lstrip("v")
            ver = f"v{int(v) + 1}"
            ws.cell(r, C["PROMPT"], text)
            ws.cell(r, C["VERSIÓN"], ver)
            log(hs, rid, ws.cell(r, C["ITEM"]).value, ver, text, cambio)
            wb.save(path)
            print(f"OK → {rid} ahora en {ver}")
            return
    raise SystemExit(f"No encontré el ID {rid}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    a = sub.add_parser("crear"); a.add_argument("path"); a.add_argument("--titulo", required=True)
    a.add_argument("--tipo", default="serie")
    b = sub.add_parser("agregar"); b.add_argument("path"); b.add_argument("desglose")
    b.add_argument("--propuesta")
    c = sub.add_parser("prompt"); c.add_argument("path"); c.add_argument("id"); c.add_argument("texto")
    c.add_argument("--cambio", default="")
    x = ap.parse_args()
    if x.cmd == "crear":
        crear(x.path, x.titulo, x.tipo)
    elif x.cmd == "agregar":
        agregar(x.path, x.desglose, x.propuesta)
    else:
        nuevo_prompt(x.path, x.id, x.texto, x.cambio)
