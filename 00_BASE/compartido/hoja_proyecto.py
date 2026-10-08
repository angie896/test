#!/usr/bin/env python3
"""
hoja_proyecto.py — la hoja de arte de un proyecto, en archivo .xlsx (modo sin conector).

Con el conector de Google Sheets, Claude hace estas mismas operaciones directo en la hoja
de Angie (ver SKILL.md). Este script sirve de respaldo y para crear la plantilla.

    python hoja_proyecto.py crear   PROYECTO.xlsx
    python hoja_proyecto.py agregar PROYECTO.xlsx desglose.json [--propuesta propuesta.json]
    python hoja_proyecto.py prompt  PROYECTO.xlsx "ITEM" "nuevo prompt" --cambio "qué cambió"

Columnas: CATEGORY | ITEM | APARECE EN | DESCRIPTION | REFERENCE | FINAL | STATUS | PROMPT
- Un ítem = una fila para todo el proyecto (no se duplica entre episodios o secuencias).
- Variantes con prompt propio = fila propia debajo: "Int. Kitchen — night".
- El prompt anterior nunca se pierde: queda en HISTORIAL PROMPTS.
- Nunca toca REFERENCE ni FINAL (las imágenes las pone Angie).
"""
import argparse
import datetime as dt
import json

from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

MONO = "Courier New"
HEAD = PatternFill("solid", fgColor="3F3F3F")
BANDS = {"SETS": "E4E9F0", "CHARACTERS": "F2F2F2", "PROPS": "E9EEE6",
         "GRAPHIC DESIGN": "F3ECE4", "VEHICLES": "EFEFEF", "MAKEUP / SFX": "F4E9EE", "VFX": "ECE9F3"}
PROMPT_FILL = PatternFill("solid", fgColor="FFF8E1")
THIN = Side(style="thin", color="A6A6A6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="top", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
FLAG = "C55A11"

COLS = [("CATEGORY", 16), ("ITEM", 30), ("APARECE EN", 22), ("DESCRIPTION", 60),
        ("REFERENCE", 24), ("FINAL", 24), ("STATUS", 15), ("PROMPT", 75)]
C = {name: i + 1 for i, (name, _) in enumerate(COLS)}
HIST = [("FECHA", 12), ("ITEM", 30), ("PROMPT", 80), ("QUÉ CAMBIÓ", 36), ("¿FUNCIONÓ?", 12)]
STATUSES = ["TO DO", "WIP", "IN REVIEW", "ADJUSTMENTS", "DONE"]
STATUS_COLORS = {"TO DO": "EDEDED", "WIP": "FFF2CC", "IN REVIEW": "DDEBF7",
                 "ADJUSTMENTS": "FCE4D6", "DONE": "E2EFDA"}
ORDER = list(BANDS)
G, H = "DESGLOSE GENERAL", "HISTORIAL PROMPTS"


def header(ws, cols):
    for i, (name, width) in enumerate(cols, start=1):
        c = ws.cell(1, i, name)
        c.fill, c.border, c.alignment = HEAD, BORDER, CENTER
        c.font = Font(name=MONO, bold=True, color="FFFFFF", size=10)
        ws.column_dimensions[c.column_letter].width = width
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "C2" if cols is COLS else "A2"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def crear(path):
    wb = Workbook()
    ws = wb.active
    ws.title = G
    header(ws, COLS)
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    L = ws.cell(1, C["STATUS"]).column_letter
    dv.add(f"{L}2:{L}3000")
    hs = wb.create_sheet(H)
    header(hs, HIST)
    dv2 = DataValidation(type="list", formula1='"sí,no,a medias"', allow_blank=True)
    hs.add_data_validation(dv2)
    dv2.add("E2:E5000")
    wb.save(path)
    print(f"OK → {path} creado")


def style_row(ws, r):
    cat = str(ws.cell(r, C["CATEGORY"]).value or "").upper()
    fill = PatternFill("solid", fgColor=BANDS.get(cat, "FFFFFF"))
    for name, col in C.items():
        c = ws.cell(r, col)
        c.border = BORDER
        c.font = Font(name=MONO, size=9, bold=name in ("CATEGORY", "ITEM"))
        c.alignment = CENTER if name in ("CATEGORY", "ITEM", "APARECE EN", "STATUS") else WRAP
        if name == "PROMPT":
            c.fill = PROMPT_FILL
        elif name == "STATUS":
            c.fill = PatternFill("solid", fgColor=STATUS_COLORS.get(str(c.value), "FFFFFF"))
        elif name not in ("REFERENCE", "FINAL"):
            c.fill = fill
    dcell = ws.cell(r, C["DESCRIPTION"])
    desc = str(dcell.value or "")
    if "⚠" in desc and not isinstance(dcell.value, CellRichText):
        # solo las líneas con ⚠ en naranja
        normal, flag = InlineFont(rFont=MONO, sz=9), InlineFont(rFont=MONO, sz=9, color=FLAG, b=True)
        ls = desc.split("\n")
        dcell.value = CellRichText(*[TextBlock(flag if l.lstrip().startswith("⚠") else normal,
                                               l + ("\n" if i < len(ls) - 1 else ""))
                                     for i, l in enumerate(ls)])
    prompt = str(ws.cell(r, C["PROMPT"]).value or "")
    lines = max(sum(1 + len(l) // 75 for l in desc.split("\n")),
                sum(1 + len(l) // 95 for l in prompt.split("\n")) if prompt else 0)
    ws.row_dimensions[r].height = max(80, min(409, lines * 12 + 12))


def find(ws, item, cat=None):
    """Fila del ítem. El mismo nombre puede existir en dos categorías (reloj físico en PROPS
    y su interfaz en GRAPHIC DESIGN), así que se compara categoría + nombre."""
    hits = [r for r in range(2, ws.max_row + 1)
            if str(ws.cell(r, C["ITEM"]).value or "").strip().lower() == item.strip().lower()
            and (cat is None or str(ws.cell(r, C["CATEGORY"]).value or "").upper() == cat.upper())]
    if len(hits) > 1:
        raise SystemExit(f'"{item}" existe en varias categorías: usa --categoria')
    return hits[0] if hits else None


def last_row_of_category(ws, cat):
    """Fila donde insertar un ítem nuevo para que quede agrupado con su categoría."""
    rank = ORDER.index(cat) if cat in ORDER else len(ORDER)
    last = 1
    for r in range(2, ws.max_row + 1):
        v = str(ws.cell(r, C["CATEGORY"]).value or "").upper()
        if not v:
            continue
        if (ORDER.index(v) if v in ORDER else len(ORDER)) <= rank:
            last = r
    return last + 1


def insert_at(ws, r, values):
    if r <= ws.max_row:
        ws.insert_rows(r)
    for k, v in values.items():
        ws.cell(r, C[k], v)


def log(hs, item, prompt, cambio):
    r = hs.max_row + 1
    for i, v in enumerate([dt.date.today().isoformat(), item, prompt, cambio, None], start=1):
        c = hs.cell(r, i, v)
        c.border, c.alignment = BORDER, WRAP if i in (3, 4) else CENTER
        c.font = Font(name=MONO, size=9)
    hs.row_dimensions[r].height = 90


def agregar(path, desglose, propuesta):
    wb = load_workbook(path)
    ws, hs = wb[G], wb[H]
    with open(desglose, encoding="utf-8") as f:
        d = json.load(f)
    prompts = {}
    if propuesta:
        with open(propuesta, encoding="utf-8") as f:
            for it in json.load(f)["items"]:
                if it.get("prompts"):
                    prompts[(it["category"].upper(), it["item"].strip().lower())] = it["prompts"]
    nuevos = actualizados = con_prompt = 0
    for sh in d["sheets"]:
        unit = sh["name"]
        for it in sh["items"]:
            cat, desc = it["category"].upper(), it.get("description", "")
            where = f'{unit}: esc. {it["scenes"]}' if it.get("scenes") else unit
            r = find(ws, it["item"], cat)
            if r:
                a = ws.cell(r, C["APARECE EN"])
                if unit not in str(a.value or ""):
                    a.value = f"{a.value}\n{where}" if a.value else where
                dc = ws.cell(r, C["DESCRIPTION"])
                if isinstance(dc.value, CellRichText):
                    dc.value = str(dc.value)
                if desc and desc not in str(dc.value or ""):
                    dc.value = f"{dc.value}\n[{unit}] {desc}" if dc.value else desc
                actualizados += 1
            else:
                r = last_row_of_category(ws, cat)
                insert_at(ws, r, {"CATEGORY": cat, "ITEM": it["item"], "APARECE EN": where,
                                  "DESCRIPTION": desc, "STATUS": "TO DO"})
                nuevos += 1
            for n, pr in enumerate(prompts.pop((cat, it["item"].strip().lower()), [])):
                if n == 0:
                    row, name = r, it["item"]
                    if ws.cell(row, C["PROMPT"]).value:
                        continue
                else:  # variante con prompt propio: fila justo debajo
                    name = f'{it["item"]} — {pr.get("label", f"v{n + 1}")}'
                    row = r + n
                    insert_at(ws, row, {"CATEGORY": cat, "ITEM": name, "APARECE EN": where,
                                        "STATUS": "TO DO"})
                ws.cell(row, C["PROMPT"], pr["text"])
                log(hs, name, pr["text"], "versión inicial")
                con_prompt += 1
    for r in range(2, ws.max_row + 1):  # insert_rows no mueve alturas ni estilos
        if ws.cell(r, C["ITEM"]).value:
            style_row(ws, r)
    wb.save(path)
    print(f"OK → {path}: {nuevos} nuevos, {actualizados} actualizados, {con_prompt} prompts")


def nuevo_prompt(path, item, text, cambio, cat=None):
    wb = load_workbook(path)
    ws, hs = wb[G], wb[H]
    r = find(ws, item, cat)
    if not r:
        raise SystemExit(f'No encontré "{item}" en la hoja')
    ws.cell(r, C["PROMPT"], text)
    log(hs, ws.cell(r, C["ITEM"]).value, text, cambio)
    wb.save(path)
    print(f'OK → prompt de "{item}" actualizado; el anterior sigue en {H}')


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    a = sub.add_parser("crear"); a.add_argument("path")
    b = sub.add_parser("agregar"); b.add_argument("path"); b.add_argument("desglose")
    b.add_argument("--propuesta")
    c = sub.add_parser("prompt"); c.add_argument("path"); c.add_argument("item"); c.add_argument("texto")
    c.add_argument("--cambio", default="")
    c.add_argument("--categoria")
    x = ap.parse_args()
    if x.cmd == "crear":
        crear(x.path)
    elif x.cmd == "agregar":
        agregar(x.path, x.desglose, x.propuesta)
    else:
        nuevo_prompt(x.path, x.item, x.texto, x.cambio, x.categoria)
